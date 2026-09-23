from __future__ import annotations

import json
from collections import Counter, defaultdict
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
DOC_DIR = ROOT / "DOCUMENTACAO_ATUAL"
SOURCE = DOC_DIR / "inventarios" / "funcoes.json"
OUTPUT = DOC_DIR / "CATALOGO_COMPLETO_DE_FUNCOES.docx"


def font(run, name: str = "Aptos", size: float | None = None, bold: bool | None = None, color: str | None = None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    return run


def shade(cell, fill: str):
    props = cell._tc.get_or_add_tcPr()
    shd = props.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        props.append(shd)
    shd.set(qn("w:fill"), fill)


def borders(cell, color: str = "D9D9D9"):
    props = cell._tc.get_or_add_tcPr()
    node = props.find(qn("w:tcBorders"))
    if node is None:
        node = OxmlElement("w:tcBorders")
        props.append(node)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = node.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            node.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), color)


def margins(cell, top=90, start=100, bottom=90, end=100):
    props = cell._tc.get_or_add_tcPr()
    tc_mar = props.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        props.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        child = tc_mar.find(qn(f"w:{name}"))
        if child is None:
            child = OxmlElement(f"w:{name}")
            tc_mar.append(child)
        child.set(qn("w:w"), str(value))
        child.set(qn("w:type"), "dxa")


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Página ")
    font(run, size=8, color="66736D")
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def display_params(value: str) -> str:
    value = value.strip()
    if not value or value == "nenhum":
        return "—"
    return value if len(value) <= 150 else value[:147] + "..."


def build():
    functions = json.loads(SOURCE.read_text(encoding="utf-8"))
    by_area = defaultdict(list)
    for item in functions:
        by_area[item["area"]].append(item)
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Cm(1.35)
    section.bottom_margin = Cm(1.35)
    section.left_margin = Cm(1.45)
    section.right_margin = Cm(1.45)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = RGBColor.from_string("26332E")
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08
    for style_name, size, before, after in (("Title", 28, 0, 14), ("Heading 1", 19, 15, 7), ("Heading 2", 13, 11, 5)):
        style = styles[style_name]
        style.font.name = "Aptos Display" if style_name != "Title" else "Georgia"
        style._element.rPr.rFonts.set(qn("w:ascii"), style.font.name)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), style.font.name)
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = style_name != "Title"
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    catalog_title = styles.add_style("Catalog Title", WD_STYLE_TYPE.PARAGRAPH)
    catalog_title.base_style = styles["Normal"]
    catalog_title.font.name = "Georgia"
    catalog_title._element.rPr.rFonts.set(qn("w:ascii"), "Georgia")
    catalog_title._element.rPr.rFonts.set(qn("w:hAnsi"), "Georgia")
    catalog_title.font.size = Pt(28)
    catalog_title.font.color.rgb = RGBColor(0, 0, 0)
    catalog_title.paragraph_format.space_after = Pt(14)
    catalog_title.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.text = "AUTOR COPILOT   ·   CATÁLOGO TÉCNICO"
    font(header.runs[0], size=8, bold=True, color="356955")
    header.paragraph_format.space_after = Pt(0)
    add_page_number(section.footer.paragraphs[0])

    title = doc.add_paragraph(style="Catalog Title")
    title.add_run("Catálogo completo de funções do Autor Copilot")
    subtitle = doc.add_paragraph()
    font(subtitle.add_run("Inventário técnico do código atualmente presente no projeto"), size=12, color="53625C")
    intro = doc.add_paragraph()
    intro.add_run("Este documento lista as funções nomeadas encontradas no código-fonte, com localização, tipo, visibilidade e parâmetros. ")
    intro.add_run("Ele serve como índice para manutenção, revisão arquitetural e entrada de novos colaboradores.")

    counts = Counter(item["area"] for item in functions)
    summary = doc.add_table(rows=1, cols=3)
    summary.alignment = WD_TABLE_ALIGNMENT.LEFT
    summary.autofit = False
    widths = [Cm(7.5), Cm(4), Cm(14.5)]
    for cell, width in zip(summary.rows[0].cells, widths):
        cell.width = width
    for cell, value in zip(summary.rows[0].cells, [f"{len(functions)} funções catalogadas", f"{len(by_area)} áreas", "Snapshot gerado diretamente do repositório"]):
        shade(cell, "EAF3EF")
        borders(cell)
        margins(cell, 150, 170, 150, 170)
        font(cell.paragraphs[0].add_run(value), size=10, bold=True, color="285A47")

    doc.add_heading("Índice por área", level=1)
    index_table = doc.add_table(rows=1, cols=2)
    index_table.alignment = WD_TABLE_ALIGNMENT.LEFT
    index_table.autofit = False
    index_table.columns[0].width = Cm(18)
    index_table.columns[1].width = Cm(4)
    for idx, text in enumerate(["Área", "Funções"]):
        cell = index_table.rows[0].cells[idx]
        shade(cell, "244B3D")
        borders(cell)
        margins(cell)
        font(cell.paragraphs[0].add_run(text), size=9, bold=True, color="FFFFFF")
    set_repeat_header(index_table.rows[0])
    for row_index, area in enumerate(sorted(counts)):
        cells = index_table.add_row().cells
        for cell in cells:
            borders(cell); margins(cell)
            if row_index % 2: shade(cell, "F4F7F5")
        font(cells[0].paragraphs[0].add_run(area), size=9)
        cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        font(cells[1].paragraphs[0].add_run(str(counts[area])), size=9, bold=True)

    doc.add_page_break()
    doc.add_heading("Como ler o catálogo", level=1)
    doc.add_paragraph("Função identifica a declaração no código. Tipo distingue funções tradicionais, arrows e versões assíncronas. Visibilidade informa se a declaração é exportada pelo módulo. O caminho e a linha permitem abrir diretamente a implementação. Parâmetros são apresentados de forma compacta; assinaturas muito longas são abreviadas no Word e permanecem completas no inventário Markdown e JSON da pasta.")

    for area in sorted(by_area):
        doc.add_heading(area, level=1)
        doc.add_paragraph(f"{len(by_area[area])} funções nomeadas nesta área.")
        grouped = defaultdict(list)
        for item in by_area[area]:
            grouped[item["file"]].append(item)
        for file_name in sorted(grouped):
            heading = doc.add_paragraph(style="Heading 2")
            font(heading.add_run(file_name), name="Aptos", size=11.5, bold=True, color="000000")
            table = doc.add_table(rows=1, cols=5)
            table.alignment = WD_TABLE_ALIGNMENT.LEFT
            table.autofit = False
            col_widths = [Cm(5.1), Cm(3.2), Cm(2), Cm(1.4), Cm(14.3)]
            labels = ["Função", "Tipo", "Visibilidade", "Linha", "Parâmetros"]
            for idx, (cell, width, label) in enumerate(zip(table.rows[0].cells, col_widths, labels)):
                cell.width = width
                shade(cell, "244B3D")
                borders(cell)
                margins(cell)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                font(cell.paragraphs[0].add_run(label), size=8.5, bold=True, color="FFFFFF")
                cell.paragraphs[0].paragraph_format.keep_with_next = True
                if idx in (2, 3): cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_repeat_header(table.rows[0])
            for row_index, item in enumerate(grouped[file_name]):
                cells = table.add_row().cells
                prevent_split(table.rows[-1])
                values = [item["name"], item["kind"], "Exportada" if item["exported"] else "Interna", str(item["line"]), display_params(item["params"])]
                for idx, (cell, width, value) in enumerate(zip(cells, col_widths, values)):
                    cell.width = width
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    borders(cell); margins(cell)
                    if row_index % 2: shade(cell, "F4F7F5")
                    run = cell.paragraphs[0].add_run(str(value))
                    font(run, name="Consolas" if idx in (0, 4) else "Aptos", size=8.2 if idx in (0, 4) else 8.5, bold=idx == 0, color="204E3C" if idx == 0 else "26332E")
                    if idx in (2, 3): cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            doc.add_paragraph().paragraph_format.space_after = Pt(1)

    doc.core_properties.title = "Catálogo completo de funções do Autor Copilot"
    doc.core_properties.subject = "Inventário técnico do projeto"
    doc.core_properties.author = "Autor Copilot"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
